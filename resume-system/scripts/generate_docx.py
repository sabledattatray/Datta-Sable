"""Build editable masters from canonical facts. Run with the bundled Python."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from styles import design_tokens as T
from templates.variants import VARIANTS

def load(name):
    return json.loads((ROOT / 'data' / (name + '.json')).read_text(encoding='utf-8'))

def link(paragraph, label, target):
    rel = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), T.BLUE); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), T.FONT); fonts.set(qn('w:hAnsi'), T.FONT); props.append(fonts)
    run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text); h.append(run); paragraph._p.append(h)

def heading(doc, title):
    p = doc.add_paragraph(title.upper(), 'Heading 1')
    borders = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    for key, value in [('val', 'single'), ('sz', '4'), ('color', T.DIVIDER), ('space', '5')]: bottom.set(qn('w:' + key), value)
    borders.append(bottom); p._p.get_or_add_pPr().append(borders)

def setup():
    d = Document(); sec = d.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(T.MARGIN_MM)
    sec.header_distance = sec.footer_distance = Mm(8)
    for name in ['Normal', 'Body Text', 'List Bullet', 'Heading 1', 'Heading 2', 'Title', 'Subtitle']:
        style = d.styles[name]; style.font.name = T.FONT; style.font.color.rgb = RGBColor.from_string(T.TEXT)
        style.font.size = Pt(T.BODY_PT)
        pf = style.paragraph_format; pf.line_spacing = 1.1; pf.space_after = Pt(5); pf.widow_control = True
    h = d.styles['Heading 1']; h.font.size = Pt(10.5); h.font.bold = True; h.font.color.rgb = RGBColor.from_string(T.NAVY)
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(8); h.paragraph_format.keep_with_next = True
    h2 = d.styles['Heading 2']; h2.font.size = Pt(10.5); h2.font.bold = True; h2.font.color.rgb = RGBColor.from_string(T.NAVY)
    h2.paragraph_format.space_before = Pt(8); h2.paragraph_format.space_after = Pt(2); h2.paragraph_format.keep_with_next = True
    b = d.styles['List Bullet']; b.paragraph_format.left_indent = Mm(3.8); b.paragraph_format.first_line_indent = Mm(-3.8); b.paragraph_format.space_after = Pt(4)
    d.core_properties.author = 'Dattatray Sable'; d.core_properties.subject = 'Professional resume'
    # Footer contains pagination only; all candidate information is in the body.
    p = sec.footer.paragraphs[0]; p.alignment = 2
    r = p.add_run('Page '); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(T.SLATE)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); p._p.append(field)
    return d

def header(doc, c, v):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(c['name']); r.bold = True; r.font.size = Pt(21); r.font.color.rgb = RGBColor.from_string(T.NAVY)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(v['headline']); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string(T.BLUE)
    p = doc.add_paragraph(v['tagline']); p.paragraph_format.space_after = Pt(8)
    for r in p.runs: r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(T.SLATE)
    p = doc.add_paragraph(c['location']); p.paragraph_format.space_after = Pt(3)
    for r in p.runs: r.font.size = Pt(9)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    link(p, c['phone'], 'tel:+918010803756'); p.add_run(' | '); link(p, c['email'], 'mailto:' + c['email']); p.add_run(' | '); link(p, 'www.dattasable.com', c['website'])
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); link(p, 'linkedin.com/in/dattasable', c['linkedin'])

def jobs(doc, experience, indices):
    for i in indices:
        job = experience[i]
        doc.add_paragraph(job['title'], 'Heading 2')
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.keep_with_next = True
        r = p.add_run(job['employer']); r.bold = True; r.font.size = Pt(9.5)
        r = p.add_run(f" | {job['location']} | {job['start']} - {job['end']}"); r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(T.SLATE)
        for text in job['bullets']: doc.add_paragraph(text, 'List Bullet')

def skills(doc, v):
    heading(doc, 'Technical Skills')
    for title, text in v['skills']:
        p = doc.add_paragraph(); p.add_run(title + ': ').bold = True; p.add_run(text)

def projects(doc, v, all_projects):
    heading(doc, 'Selected Technical Projects')
    for key in v['projects']:
        item = next(x for x in all_projects if x['id'] == key)
        doc.add_paragraph(item['title'] + ' | Independent portfolio', 'Heading 2')
        p = doc.add_paragraph(item['description']); p.paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
        link(p, item['url'].replace('https://', ''), item['url'])

def education(doc, c):
    heading(doc, 'Education'); p = doc.add_paragraph(); p.add_run(c['education']['degree']).bold = True
    p.add_run(f" | {c['education']['year']}")
    doc.add_paragraph(c['education']['university'])

def main():
    c, exp, projs = load('candidate'), load('experience'), load('projects')
    if not c['dbs_end_date'] or exp[0]['end'] != c['dbs_end_date']: raise SystemExit('Confirm DBS end date before final generation.')
    out = ROOT / 'output'; out.mkdir(exist_ok=True)
    for key, v in VARIANTS.items():
        selected = next((arg.split('=', 1)[1] for arg in sys.argv if arg.startswith('--profile=')), None)
        if selected and key != selected: continue
        doc = setup(); doc.core_properties.title = c['name'] + ' - ' + v['headline']; header(doc, c, v)
        heading(doc, 'Professional Summary'); doc.add_paragraph(v['summary'])
        if key in ('master', 'power_bi_analyst'):
            skills(doc, v); projects(doc, v, projs); doc.add_page_break()
            heading(doc, 'Professional Experience'); jobs(doc, exp, [0, 1, 2, 3]); education(doc, c)
        else:
            heading(doc, 'Core Competencies'); doc.add_paragraph(v['competencies'])
            heading(doc, 'Professional Experience'); jobs(doc, exp, [0, 1]); doc.add_page_break()
            heading(doc, 'Professional Experience (continued)'); jobs(doc, exp, [2, 3])
            skills(doc, v); projects(doc, v, projs); education(doc, c)
        target = out / (v['filename'] + '.docx'); doc.save(target); print(target)

if __name__ == '__main__': main()
