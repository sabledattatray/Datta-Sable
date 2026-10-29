"""Check canonical text, document order, PDF links, page geometry and duplication."""
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import xml.etree.ElementTree as ET
import pdfplumber
from pypdf import PdfReader
import pypdfium2 as pdfium
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from templates.variants import VARIANTS

def normalized(text):
    return re.sub(r'\s+', '', text).replace('\u00ad', '').replace('\u2011', '-').replace('\ufb01', 'fi').replace('\ufb02', 'fl')

def validate():
    jobs = json.loads((ROOT / 'data/experience.json').read_text())
    results = []
    for key, variant in VARIANTS.items():
        stem = variant['filename']; path = ROOT / 'output' / (stem + '.pdf')
        reader = PdfReader(path); docx_path = path.with_suffix('.docx'); document = Document(docx_path)
        assert len(reader.pages) == 2, f'{stem}: expected two pages'
        text = '\n'.join(p.extract_text() for p in reader.pages)
        normal = normalized(text)
        assert text.count('DATTATRAY SABLE') == 1
        assert text.count('resume@dattasable.com') == 1
        assert 'gmail.com' not in text and 'Present' not in text and 'END DATE' not in text
        assert '2005' in text and 'September 2026' in text and 'CASCO' not in text
        for job in jobs:
            assert text.count(job['employer']) == 1
            assert normalized(f"{job['start']} - {job['end']}") in normal
            for bullet in job['bullets']: assert normal.count(normalized(bullet)) == 1, f'{stem}: missing or duplicate bullet'
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for paragraph in paragraphs:
            assert normalized(paragraph) in normal, f'{stem}: DOCX/PDF mismatch: {paragraph}'
        meaningful = [normalized(p) for p in paragraphs if len(p) > 90]
        assert all(n == 1 for n in Counter(meaningful).values()), f'{stem}: duplicated paragraphs'
        expected_headings = ['PROFESSIONAL SUMMARY', 'TECHNICAL SKILLS', 'SELECTED TECHNICAL PROJECTS', 'PROFESSIONAL EXPERIENCE', 'EDUCATION'] if key == 'power_bi_analyst' else ['PROFESSIONAL SUMMARY', 'CORE COMPETENCIES', 'PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS', 'SELECTED TECHNICAL PROJECTS', 'EDUCATION']
        offsets = [text.index(h) for h in expected_headings]; assert offsets == sorted(offsets)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with ZipFile(docx_path) as z:
            root = ET.fromstring(z.read('word/document.xml'))
            assert not root.findall('.//w:tbl', ns) and not root.findall('.//w:txbxContent', ns)
            assert not root.findall('.//w:vanish', ns)
        links = [str(a.get_object().get('/A', {}).get('/URI', '')) for p in reader.pages for a in p.get('/Annots', [])]
        assert 'mailto:resume@dattasable.com' in links
        assert 'https://www.dattasable.com' in [url.rstrip('/') for url in links]
        assert 'https://www.linkedin.com/in/dattasable/' in links
        projects = json.loads((ROOT / 'data/projects.json').read_text())
        for p in projects:
            if p['id'] in variant['projects']: assert p['url'].rstrip('/') in [url.rstrip('/') for url in links], f"Missing PDF link: {p['url']}"
        page_stats = []
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                assert abs(page.width - 595.28) < 1 and abs(page.height - 841.89) < 1
                chars = [c for c in page.chars if c['text'].strip()]
                assert all(40 <= c['x0'] < c['x1'] <= page.width - 40 for c in chars)
                assert all(40 <= c['top'] < c['bottom'] <= page.height - 12 for c in chars)
                duplicates = Counter((round(c['x0'], 2), round(c['top'], 2), c['text']) for c in chars)
                assert max(duplicates.values()) == 1, f'{stem}: duplicate text overlays on page {i+1}'
                # Separate words must not occupy the same area on the same baseline.
                words = page.extract_words()
                for a, b in zip(words, words[1:]):
                    if abs(a['top'] - b['top']) < 1:
                        assert b['x0'] >= a['x1'] - .5, f'{stem}: overlapping words'
                page_stats.append({'page': i + 1, 'words': len(words), 'min_font_pt': min(c['size'] for c in chars)})
        preview = ROOT / 'previews' / stem; preview.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(pdfium.PdfDocument(path)):
            page.render(scale=1.5).to_pil().save(preview / f'page-{i+1}.png')
        (preview / 'extracted.txt').write_text(text, encoding='utf-8')
        results.append({'variant': key, 'file': stem, 'pages': 2, 'links': links, 'checks': 'PASS', 'page_stats': page_stats})
    (ROOT / 'validation-results.json').write_text(json.dumps(results, indent=2) + '\n')
    print(json.dumps(results, indent=2))

if __name__ == '__main__': validate()
